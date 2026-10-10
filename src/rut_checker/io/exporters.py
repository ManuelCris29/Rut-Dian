r"""Exportador de resultados: Excel (hoja `resultado`) y JSON de respaldo (ARQUITECTURA §5.3 y §7).

Reglas:
- Solo se serializan las columnas de SPEC §5, en ese orden; nunca `model_dump()` completo.
- Toda celda de texto pasa por `_sanitize` (caracteres que el XML/UTF-8 no admiten -> U+FFFD,
  sin eliminar dígitos) y luego por `neutralize_cell` (Obs-9): si empieza por `= + - @ \t \r`
  se prefija con un apóstrofo visible que forma parte del valor (sobrevive a CSV/JSON). Se
  escribe siempre como texto; la columna `identificacion` además con formato `@`.
- La nota del área (`raw_note`) de filas `ENTRADA_INVALIDA` nunca se exporta (minimización).
- `fecha_consulta` se convierte a `America/Bogota` (ISO 8601 con desfase).
- Los archivos solo se escriben en `output_dir` (lo resuelve quien llama desde `RUT_OUTPUT_DIR`;
  este módulo no importa `infra`), de forma atómica: temporal en el mismo directorio y
  `os.replace`; nunca se sobrescribe ni se borra un archivo previo (si el destino existe se
  trata como fallo y se pasa al siguiente destino).
- Fallo al escribir: un reintento con `<nombre>_1.xlsx`; si también falla, JSON de respaldo
  `<nombre>.json` (resultado `JSON_BACKUP`: la CLI sale con código 3). Solo si fallan los tres
  se lanza `ExportError` genérico. Ningún mensaje incluye valores, rutas ni texto de terceros.
- Este módulo no registra en logs.
"""

import contextlib
import json
import os
import re
from collections.abc import Callable
from dataclasses import dataclass
from enum import StrEnum
from pathlib import Path
from typing import Final
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.utils.exceptions import IllegalCharacterError

from rut_checker.domain.errors import ExportError, UnsafePathError, UnsafePathReason
from rut_checker.domain.models import QueryOutcome

COLUMNS: tuple[str, ...] = (
    "identificacion",
    "dv",
    "nombre_o_razon_social",
    "estado_rut",
    "fecha_consulta",
    "fuente",
    "estado_consulta",
    "observacion",
)
SHEET_NAME = "resultado"
_BOGOTA: Final = ZoneInfo("America/Bogota")
_DANGEROUS_PREFIXES = ("=", "+", "-", "@", "\t", "\r")
_QUOTE = "'"
_REPLACEMENT = "�"
# Controles C0 (salvo \t, \n, \r), surrogates sueltos y los no-caracteres U+FFFE/U+FFFF:
# el XML del .xlsx no los admite y un surrogate suelto no se codifica en UTF-8 (JSON).
_ILLEGAL_TEXT = re.compile("[\\x00-\\x08\\x0b\\x0c\\x0e-\\x1f\\ud800-\\udfff\\ufffe\\uffff]")
_Row = tuple[str, ...]
_Writer = Callable[[list[_Row], Path], None]


class ExportKind(StrEnum):
    """Qué se escribió realmente; la CLI sale con código 3 si fue `JSON_BACKUP`."""

    XLSX = "XLSX"
    XLSX_ALTERNATIVE = "XLSX_ALTERNATIVE"
    JSON_BACKUP = "JSON_BACKUP"


@dataclass(frozen=True)
class ExportReport:
    """Resultado de `export_excel`: ruta del archivo escrito y su tipo."""

    path: Path
    kind: ExportKind


def _sanitize(value: str) -> str:
    """Reemplaza por U+FFFD los caracteres no representables (nunca los elimina)."""
    return _ILLEGAL_TEXT.sub(_REPLACEMENT, value)


def neutralize_cell(value: str) -> str:
    """Prefija con apóstrofo el texto que Excel podría interpretar como fórmula (Obs-9)."""
    if value.startswith(_DANGEROUS_PREFIXES):
        return _QUOTE + value
    return value


def _cell_text(value: str) -> str:
    return neutralize_cell(_sanitize(value))


def _raw_text(value: str | int | float | None) -> str:
    """Texto de un valor crudo de la celda; un float entero se muestra sin `.0`."""
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def _identification_and_dv(outcome: QueryOutcome) -> tuple[str, str]:
    if outcome.identification is not None:
        return outcome.identification.identification, outcome.identification.dv or ""
    return _raw_text(outcome.row.raw_identification), outcome.row.raw_dv or ""


def _build_row(outcome: QueryOutcome) -> _Row:
    identification, dv = _identification_and_dv(outcome)
    result = outcome.result
    queried_at = outcome.queried_at.astimezone(_BOGOTA)
    values = (
        identification,
        dv,
        result.name,
        result.estado_rut,
        queried_at.isoformat(timespec="seconds"),
        result.source.value,
        result.status.value,
        result.observation,
    )
    return tuple(_cell_text(value) for value in values)


def _write_workbook(rows: list[_Row], path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    if sheet is None:  # pragma: no cover - un Workbook nuevo siempre trae hoja activa
        raise ExportError
    sheet.title = SHEET_NAME
    for col, name in enumerate(COLUMNS, start=1):
        sheet.cell(row=1, column=col, value=name)
    for row_number, values in enumerate(rows, start=2):
        for col, text in enumerate(values, start=1):
            cell = sheet.cell(row=row_number, column=col)
            # Orden obligatorio: primero asignar el valor (openpyxl infiere 'f' si empieza
            # por '=') y después forzar data_type='s'.
            cell.value = text
            cell.data_type = "s"
            if col == 1:
                cell.number_format = "@"
    workbook.save(path)


def _write_json(rows: list[_Row], path: Path) -> None:
    records = [dict(zip(COLUMNS, values, strict=True)) for values in rows]
    path.write_text(json.dumps(records, ensure_ascii=False, indent=2), encoding="utf-8")


def _atomic(path: Path, writer: _Writer, rows: list[_Row]) -> bool:
    """Escribe en un temporal del mismo directorio y lo mueve con `os.replace`.

    Devuelve False si falla (sin propagar el texto de la causa). Si el destino ya existe se
    trata como fallo: no se sobrescribe ni se borra nada previo. El temporal se borra siempre.
    """
    if path.exists():
        return False
    temp = path.with_name(f".{path.stem}.tmp{path.suffix}")
    try:
        writer(rows, temp)
        os.replace(temp, path)
    except (OSError, IllegalCharacterError, UnicodeError):
        return False
    finally:
        with contextlib.suppress(OSError):
            temp.unlink(missing_ok=True)
    return True


def _check_name(name: str) -> None:
    if not name or name in {".", ".."} or Path(name).name != name or "\\" in name:
        raise UnsafePathError(UnsafePathReason.INVALID_PATH)


def export_json_backup(outcomes: list[QueryOutcome], output_dir: Path, name: str) -> Path:
    """Escribe el JSON de respaldo `<output_dir>/<name>.json` solo con las columnas de SPEC §5.

    Args:
        outcomes: Desenlaces en el orden original de la entrada.
        output_dir: Directorio de salida (el de `RUT_OUTPUT_DIR`); debe existir.
        name: Nombre base sin extensión ni directorios (p. ej. `resultado_<run_id>`).

    Raises:
        UnsafePathError: `name` vacío o con separadores/`..`.
        ExportError: Si no se puede escribir o el destino ya existe (mensaje genérico).
    """
    _check_name(name)
    path = output_dir / f"{name}.json"
    if not _atomic(path, _write_json, [_build_row(outcome) for outcome in outcomes]):
        raise ExportError from None
    return path


def export_excel(outcomes: list[QueryOutcome], path: Path, *, output_dir: Path) -> ExportReport:
    """Escribe el Excel `resultado` con las columnas de SPEC §5.

    Si la escritura falla (p. ej. archivo abierto en Excel en Windows) o el destino ya existe,
    reintenta una vez con `<nombre>_1.xlsx` en el mismo directorio; si también falla escribe
    `<nombre>.json` en `output_dir`. Nunca sobrescribe ni borra un archivo previo.

    Args:
        outcomes: Desenlaces en el orden original de la entrada (una fila por fila de entrada).
        path: Ruta destino, `output_dir/resultado_<run_id>.xlsx` (sufijo `.xlsx`, sin distinguir
            mayúsculas).
        output_dir: Directorio de salida; `path` debe estar directamente dentro de él.

    Returns:
        `ExportReport` con la ruta escrita y su tipo (`XLSX`, `XLSX_ALTERNATIVE` o
        `JSON_BACKUP`; este último implica código de salida 3 en la CLI).

    Raises:
        UnsafePathError: Ruta fuera de `output_dir` o sufijo distinto de `.xlsx`.
        ExportError: Fallaron los tres destinos (mensaje genérico).
    """
    if path.resolve().parent != output_dir.resolve():
        raise UnsafePathError(UnsafePathReason.OUTSIDE_OUTPUT_DIR)
    if path.suffix.lower() != ".xlsx":
        raise UnsafePathError(UnsafePathReason.INVALID_PATH)
    rows = [_build_row(outcome) for outcome in outcomes]
    alternative = path.with_name(f"{path.stem}_1{path.suffix}")
    if _atomic(path, _write_workbook, rows):
        return ExportReport(path, ExportKind.XLSX)
    if _atomic(alternative, _write_workbook, rows):
        return ExportReport(alternative, ExportKind.XLSX_ALTERNATIVE)
    backup = output_dir / f"{path.stem}.json"
    if _atomic(backup, _write_json, rows):
        return ExportReport(backup, ExportKind.JSON_BACKUP)
    raise ExportError from None
