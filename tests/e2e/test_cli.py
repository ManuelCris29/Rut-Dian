"""Pruebas e2e de la CLI (T-40) con `CliRunner`, sin red ni navegador.

Cada prueba corre en un proyecto temporal (`workspace`): cwd propio, `.gitignore` con
`output/`, `logs/` y `data/`, sin `.env` y sin esperas reales. Solo se usan identificaciones
de la lista blanca ficticia.
"""

import json
import logging
import os
import re
import subprocess
import sys
from collections.abc import Iterator
from datetime import datetime, timedelta
from pathlib import Path
from zoneinfo import ZoneInfo

import pytest
from openpyxl import load_workbook
from typer.testing import CliRunner, Result

from rut_checker import cli, cli_support
from rut_checker.adapters.console_operator import ConsoleOperator
from rut_checker.adapters.contingency_source import ContingencySource
from rut_checker.adapters.document_copy_source import DocumentCopySource
from rut_checker.application.batch_service import BatchQueryService
from rut_checker.application.report import StopReason
from rut_checker.application.retry_policy import BackoffRetryPolicy
from rut_checker.cli import app, build_service
from rut_checker.domain.errors import (
    ExportError,
    PortalUnavailable,
    RutCheckerError,
    SourceNotEnabledReason,
    StructureChanged,
    UnsafePathError,
    UnsafePathReason,
)
from rut_checker.domain.models import IdentificationInput, RutResult
from rut_checker.domain.status import QueryStatus, SourceKind
from rut_checker.infra.logging import ChallengeLogSink, close_run_logging
from rut_checker.infra.paths import assert_safe_dir, ensure_private_dir
from rut_checker.infra.settings import Settings
from rut_checker.io.exporters import ExportKind, ExportReport

runner = CliRunner()
BOGOTA = ZoneInfo("America/Bogota")
ID_1, ID_2, ID_3 = "1000000001", "1000000002", "1000000003"
RUN_ID = re.compile(r"[0-9]{8}-[0-9]{6}-[0-9a-f]{4}")
APPROVAL = "APR-20261001-1"
TEST_RUN_ID = "20261010-120000-abcd"
HEADER = "identificacion,fecha_copia_rut,responsable_validacion\n"
PORTAL_FACTORY = "rut_checker.adapters.dian_portal.factory.create_portal_source"
PORTAL_GATE = "rut_checker.adapters.dian_portal.selectors.ensure_portal_enabled"

type Responses = list[RutResult | type[RutCheckerError]]


class RecordingSleeper:
    """Sustituye a `SystemSleeper`: registra las esperas sin dormir."""

    instances: list["RecordingSleeper"] = []  # noqa: RUF012  (registro de la prueba)

    def __init__(self) -> None:
        self.calls: list[float] = []
        RecordingSleeper.instances.append(self)

    def sleep(self, seconds: float) -> None:
        self.calls.append(seconds)


class PortalDouble:
    """Doble del portal: fuente con ritmo que responde según un guion."""

    kind = SourceKind.PORTAL_DIAN_C1
    paced = True

    def __init__(self, responses: Responses) -> None:
        self._responses = list(responses)
        self.calls = 0
        self.closed = 0

    def query(self, ident: IdentificationInput) -> RutResult:
        self.calls += 1
        item = self._responses.pop(0) if len(self._responses) > 1 else self._responses[0]
        if isinstance(item, RutResult):
            return item
        raise item()

    def close(self) -> None:
        self.closed += 1


PORTAL_OK = RutResult(
    status=QueryStatus.EXITOSA,
    source=SourceKind.PORTAL_DIAN_C1,
    estado_rut="ACTIVO",
    name="PERSONA FICTICIA UNO",
)


@pytest.fixture
def workspace(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> Iterator[Path]:
    """Proyecto temporal: cwd propio, carpetas por defecto ignoradas por git y sin esperas."""
    (tmp_path / ".gitignore").write_text("output/\nlogs/\ndata/\n", encoding="utf-8")
    monkeypatch.chdir(tmp_path)
    RecordingSleeper.instances = []
    monkeypatch.setattr(cli, "SystemSleeper", RecordingSleeper)
    root = logging.getLogger()
    saved = list(root.handlers)  # la CLI añade un NullHandler al root; no debe filtrarse
    yield tmp_path
    root.handlers[:] = saved


def _recent_copy_date(days: int = 10) -> str:
    return (datetime.now(BOGOTA).date() - timedelta(days=days)).isoformat()


def _write_batch(workspace: Path, lines: list[str], name: str = "lote.csv") -> Path:
    path = workspace / name
    path.write_text(HEADER + "\n".join(lines) + "\n", encoding="utf-8")
    return path


def _lote(path: Path, *extra: str) -> Result:
    return runner.invoke(app, ["lote", str(path), *extra])


def _rows(workspace: Path) -> list[tuple[str, ...]]:
    (excel,) = sorted((workspace / "output").glob("resultado_*.xlsx"))
    sheet = load_workbook(excel)["resultado"]
    return [tuple(str(cell or "") for cell in row) for row in sheet.iter_rows(values_only=True)]


def _use_fake(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setenv("RUT_SOURCE", "fake")
    monkeypatch.setenv("RUT_ALLOW_FAKE", "true")


def _enable_portal(monkeypatch: pytest.MonkeyPatch, portal: PortalDouble) -> None:
    """Habilita el portal con dobles: ni selectores reales ni navegador."""
    monkeypatch.setenv("RUT_SOURCE", "dian_portal")
    monkeypatch.setenv("RUT_PORTAL_ENABLED", "true")
    monkeypatch.setenv("RUT_PORTAL_APPROVAL_REF", APPROVAL)
    monkeypatch.setattr(PORTAL_GATE, lambda _settings: object())
    monkeypatch.setattr(PORTAL_FACTORY, lambda *_args, **_kwargs: portal)


def _built_service() -> BatchQueryService:
    """`build_service` real; cierra el log de inmediato (el archivo no se usa en estas pruebas)."""
    service = build_service(Settings(), run_id=TEST_RUN_ID)
    close_run_logging(TEST_RUN_ID)
    return service


def _tty(monkeypatch: pytest.MonkeyPatch, *, interactive: bool = True) -> None:
    monkeypatch.setattr(cli, "stdin_is_tty", lambda: interactive)


def _log_events(workspace: Path) -> list[dict[str, object]]:
    (log,) = (workspace / "logs").glob("run_*.jsonl")
    return [json.loads(line) for line in log.read_text(encoding="utf-8").splitlines()]


def _no_leak(result: Result, workspace: Path, *ids: str) -> None:
    texts = [result.output, result.stderr]
    texts += [p.read_text(encoding="utf-8") for p in (workspace / "logs").glob("*.jsonl")]
    for ident in ids:
        assert not any(ident in text for text in texts), "identificación completa filtrada"


# --- configuración y llaves ----------------------------------------------------------------


def test_dian_portal_disabled_exits_2_without_browser_or_network(
    workspace: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """G11: portal deshabilitado -> código 2, cita (a)-(d), sin navegador ni red ni archivos."""
    monkeypatch.setenv("RUT_SOURCE", "dian_portal")

    def _forbidden(*_args: object, **_kwargs: object) -> None:
        pytest.fail("no se debe construir el adaptador del portal (navegador/red)")

    monkeypatch.setattr(PORTAL_FACTORY, _forbidden)
    result = _lote(_write_batch(workspace, [ID_1]))
    assert result.exit_code == 2
    assert "(a) a (d)" in result.stderr
    assert "deshabilitada" in result.stderr
    assert not (workspace / "logs").exists()
    assert not (workspace / "output").exists()


def test_dian_portal_with_unvalidated_selectors_exits_2(
    workspace: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """Llave 3 real: el `config/selectors.yaml` del repo no está validado -> código 2."""
    monkeypatch.setenv("RUT_SOURCE", "dian_portal")
    monkeypatch.setenv("RUT_PORTAL_ENABLED", "true")
    monkeypatch.setenv("RUT_PORTAL_APPROVAL_REF", APPROVAL)
    result = _lote(_write_batch(workspace, [ID_1]))
    assert result.exit_code == 2
    assert "(a) a (d)" in result.stderr
    assert not (workspace / "logs").exists()


def test_dian_portal_requires_the_enabled_portal_token_exits_2(
    workspace: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """`DianPortalSource` (T-49) solo acepta `EnabledPortal`: sin él la CLI no degrada."""
    monkeypatch.setenv("RUT_SOURCE", "dian_portal")
    monkeypatch.setenv("RUT_PORTAL_ENABLED", "true")
    monkeypatch.setenv("RUT_PORTAL_APPROVAL_REF", APPROVAL)
    monkeypatch.setattr(PORTAL_GATE, lambda _settings: object())
    result = _lote(_write_batch(workspace, [ID_1]))
    assert result.exit_code == 2
    assert SourceNotEnabledReason.NOT_ENABLED_PORTAL.value in result.stderr
    assert not (workspace / "logs").exists()


def test_fake_without_allow_flag_exits_2(workspace: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setenv("RUT_SOURCE", "fake")
    result = _lote(_write_batch(workspace, [ID_1]))
    assert result.exit_code == 2
    assert "pruebas" in result.stderr
    assert "DATOS DE PRUEBA" not in result.stderr
    assert not (workspace / "output").exists()


def test_invalid_settings_exit_2_listing_names_never_values(
    workspace: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setenv("RUT_PORTAL_URL", "https://evil.example.com/SECRETO-123")
    monkeypatch.setenv("RUT_MAX_RETRIES", "valor-raro")
    result = _lote(_write_batch(workspace, [ID_1]))
    assert result.exit_code == 2
    assert "RUT_PORTAL_URL" in result.stderr
    assert "RUT_MAX_RETRIES" in result.stderr
    for secret in ("evil", "SECRETO", "valor-raro"):
        assert secret not in result.output
        assert secret not in result.stderr


def test_incoherent_settings_exit_2_without_values(
    workspace: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setenv("RUT_RETRY_BACKOFF_S", "1,2")  # < 3 s: error de coherencia (modelo)
    result = _lote(_write_batch(workspace, [ID_1]))
    assert result.exit_code == 2
    assert "Configuración inválida" in result.stderr


def test_env_file_notice_shows_path_without_values(
    workspace: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    env_file = workspace / ".env"
    env_file.write_text("RUT_SOURCE=contingency\nRUT_PORTAL_APPROVAL_REF=APR-20261001-9\n")
    monkeypatch.setitem(Settings.model_config, "env_file", str(env_file))
    result = _lote(_write_batch(workspace, [ID_1]))
    assert result.exit_code == 0, result.stderr
    assert str(env_file) in result.stderr
    assert "APR-20261001-9" not in result.output
    assert "APR-20261001-9" not in result.stderr


# --- cableado de fuentes -------------------------------------------------------------------


def test_lote_fake_writes_excel_and_exits_0(
    workspace: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    _use_fake(monkeypatch)
    result = _lote(_write_batch(workspace, [ID_1, ID_2, "12A45", ID_1]))
    assert result.exit_code == 0, result.stderr
    assert "DATOS DE PRUEBA" in result.stderr
    rows = _rows(workspace)
    assert [row[6] for row in rows[1:]] == ["EXITOSA", "EXITOSA", "ENTRADA_INVALIDA", "EXITOSA"]
    assert "EXITOSA: 3" in result.output
    assert "ENTRADA_INVALIDA: 1" in result.output
    assert "resultado_" in result.output
    _no_leak(result, workspace, ID_1, ID_2)


def test_contingency_wires_document_copy_with_contingency_fallback(workspace: Path) -> None:
    """G9: con `RUT_SOURCE=contingency` una copia vigente da VALIDADA_COPIA_RUT."""
    batch = _write_batch(workspace, [f"{ID_1},{_recent_copy_date()},USR-001", f"{ID_2},,"])
    result = _lote(batch)
    assert result.exit_code == 0, result.stderr
    rows = _rows(workspace)
    assert rows[1][6] == "VALIDADA_COPIA_RUT"
    assert rows[1][5] == "COPIA_RUT_TERCERO"
    assert rows[1][2] == rows[1][3] == ""  # D-3: sin nombre ni estado
    assert rows[2][6] == "PENDIENTE_VALIDACION_MANUAL"
    assert "DATOS DE PRUEBA" not in result.stderr
    _no_leak(result, workspace, ID_1, ID_2)


def test_contingency_fallback_is_document_copy_over_contingency(workspace: Path) -> None:
    """Decisión de T-40: tras una parada, el respaldo sigue validando contra la copia."""
    service = _built_service()
    primary = service._primary
    fallback = service._fallback
    assert isinstance(primary, DocumentCopySource)
    assert isinstance(fallback, DocumentCopySource)
    assert primary is not fallback
    assert isinstance(primary._fallback, ContingencySource)
    assert isinstance(fallback._fallback, ContingencySource)
    assert isinstance(service._operator, ConsoleOperator)


def test_dian_portal_wires_document_copy_with_contingency_fallback(
    workspace: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """ADR-0002 §3: tras la parada las filas con copia vigente quedan VALIDADA_COPIA_RUT."""
    portal = PortalDouble([StructureChanged])
    _enable_portal(monkeypatch, portal)
    batch = _write_batch(
        workspace, [f"{ID_1},,", f"{ID_2},{_recent_copy_date()},USR-001", f"{ID_3},,"]
    )
    result = _lote(batch)
    assert result.exit_code == 1, result.stderr  # lote detenido
    rows = _rows(workspace)
    assert [row[6] for row in rows[1:]] == [
        "ESTRUCTURA_CAMBIADA",
        "VALIDADA_COPIA_RUT",
        "PENDIENTE_VALIDACION_MANUAL",
    ]
    assert portal.calls == 1
    assert portal.closed == 1
    assert StopReason.STRUCTURE_CHANGED.value in result.output
    assert "detenido" in result.output.lower()
    _no_leak(result, workspace, ID_1, ID_2, ID_3)


def test_dian_portal_uses_backoff_retry_policy_from_settings(
    workspace: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    portal = PortalDouble([PortalUnavailable, PORTAL_OK])
    _enable_portal(monkeypatch, portal)
    monkeypatch.setenv("RUT_RETRY_BACKOFF_S", "7,9")
    result = _lote(_write_batch(workspace, [ID_1]))
    assert result.exit_code == 0, result.stderr
    assert portal.calls == 2
    (sleeper,) = RecordingSleeper.instances
    assert 7.0 in sleeper.calls  # max(RUT_MIN_DELAY=5, backoff[0]=7)
    assert _rows(workspace)[1][6] == "EXITOSA"


def test_dian_portal_wiring_passes_gate_result_settings_and_operator(
    workspace: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    received: list[object] = []
    enabled = object()
    portal = PortalDouble([PORTAL_OK])
    _enable_portal(monkeypatch, portal)
    monkeypatch.setattr(PORTAL_GATE, lambda _settings: enabled)

    def _factory(*args: object, **_kwargs: object) -> PortalDouble:
        received.extend(args)
        return portal

    monkeypatch.setattr(PORTAL_FACTORY, _factory)
    service = _built_service()
    assert received[0] is enabled
    assert isinstance(received[1], Settings)
    assert isinstance(received[2], ConsoleOperator)
    assert isinstance(received[4], ChallengeLogSink)  # recolector de desafíos ya enlazado al log
    assert received[4]._logger is not None
    assert isinstance(service._retry_policy, BackoffRetryPolicy)


def test_contingency_does_not_import_playwright_or_yaml(tmp_path: Path) -> None:
    code = (
        "import sys\n"
        "from rut_checker.cli import build_service\n"
        "from rut_checker.infra.settings import Settings\n"
        "build_service(Settings(_env_file=None))\n"
        "loaded = {name.split('.')[0] for name in sys.modules}\n"
        "bad = loaded & {'playwright', 'yaml'}\n"
        "assert not bad, bad\n"
    )
    (tmp_path / ".gitignore").write_text("output/\nlogs/\ndata/\n", encoding="utf-8")
    completed = subprocess.run(  # noqa: S603  (intérprete propio, código fijo)
        [sys.executable, "-I", "-c", code],
        cwd=tmp_path,
        capture_output=True,
        text=True,
        timeout=120,
        check=False,
    )
    assert completed.returncode == 0, completed.stderr


# --- consultar -----------------------------------------------------------------------------


def test_consultar_reads_id_from_prompt_and_prints_masked(
    workspace: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    _use_fake(monkeypatch)
    _tty(monkeypatch)
    result = runner.invoke(app, ["consultar"], input=f"{ID_1}\n")
    assert result.exit_code == 0, result.stderr
    assert "******0001" in result.output
    _no_leak(result, workspace, ID_1)
    assert _rows(workspace)[1][0] == ID_1  # el Excel sí lleva el número (perímetro autorizado)


def test_consultar_accepts_argument_and_copy_options(workspace: Path) -> None:
    args = ["consultar", ID_2, "--fecha-copia", _recent_copy_date(), "--responsable", "USR-001"]
    result = runner.invoke(app, args)
    assert result.exit_code == 0, result.stderr
    assert "VALIDADA_COPIA_RUT" in result.output
    assert "******0002" in result.output
    _no_leak(result, workspace, ID_2)


def test_consultar_without_number_and_without_terminal_exits_2(
    workspace: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    _tty(monkeypatch, interactive=False)
    result = runner.invoke(app, ["consultar"], input=f"{ID_1}\n")
    assert result.exit_code == 2
    assert "requiere una terminal interactiva" in result.stderr
    assert not (workspace / "logs").exists()
    assert not (workspace / "output").exists()
    _no_leak(result, workspace, ID_1)


def test_consultar_prompt_eof_with_terminal_aborts_with_code_1_and_no_files(
    workspace: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    _tty(monkeypatch)
    result = runner.invoke(app, ["consultar"], input="")
    assert result.exit_code == 1
    assert not (workspace / "logs").exists()
    assert not (workspace / "output").exists()


def test_consultar_argument_warns_about_terminal_history(workspace: Path) -> None:
    result = runner.invoke(app, ["consultar", ID_1])
    assert "quedó en el historial del terminal" in result.stderr
    assert "rut-checker consultar` sin argumento" in result.stderr


def test_consultar_prompt_does_not_warn_about_history(
    workspace: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    _tty(monkeypatch)
    result = runner.invoke(app, ["consultar"], input=f"{ID_1}\n")
    assert "historial" not in result.stderr


def test_consultar_invalid_number_exits_2_but_keeps_the_excel(workspace: Path) -> None:
    result = runner.invoke(app, ["consultar", "12A45"])
    assert result.exit_code == 2
    assert "ENTRADA_INVALIDA" in result.output
    assert _rows(workspace)[1][6] == "ENTRADA_INVALIDA"


def test_consultar_invalid_copy_date_is_not_echoed(workspace: Path) -> None:
    result = runner.invoke(app, ["consultar", ID_1, "--fecha-copia", "no-es-fecha-SECRETA"])
    assert result.exit_code == 2  # la fila queda ENTRADA_INVALIDA; el valor no se repite
    assert "SECRETA" not in result.output
    assert "SECRETA" not in result.stderr


# --- errores de uso de Click: no repiten valores (M-1) -----------------------------------


@pytest.mark.parametrize(
    "args",
    [
        ["consultar", ID_1, ID_2],
        ["consultar", f"--{ID_2}"],
        ["consultar", "--dv"],
        ["lote", "lote.csv", ID_2],
        [ID_2],
    ],
)
def test_usage_errors_exit_2_without_echoing_values(workspace: Path, args: list[str]) -> None:
    result = runner.invoke(app, args)
    assert result.exit_code == 2
    assert ID_1 not in result.output
    assert ID_2 not in result.output
    assert ID_2 not in result.stderr
    assert not (workspace / "logs").exists()


def test_missing_required_argument_keeps_its_name(workspace: Path) -> None:
    result = runner.invoke(app, ["lote"])
    assert result.exit_code == 2
    assert "ARCHIVO" in result.output


# --- salida, rutas y códigos ---------------------------------------------------------------


def test_json_backup_exits_with_code_3(workspace: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    backup = workspace / "output" / "resultado_x.json"
    monkeypatch.setattr(
        cli, "export_excel", lambda *_a, **_k: ExportReport(backup, ExportKind.JSON_BACKUP)
    )
    result = _lote(_write_batch(workspace, [ID_1]))
    assert result.exit_code == 3
    assert "respaldo" in result.stderr.lower()
    assert "JSON" in result.stderr


def test_export_error_exits_with_code_3(workspace: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    def _fail(*_a: object, **_k: object) -> ExportReport:
        raise ExportError

    monkeypatch.setattr(cli, "export_excel", _fail)
    result = _lote(_write_batch(workspace, [ID_1]))
    assert result.exit_code == 3
    assert "No se pudo exportar" in result.stderr


def test_alternative_excel_is_reported_with_code_0(
    workspace: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    alt = workspace / "output" / "resultado_x_1.xlsx"
    monkeypatch.setattr(
        cli, "export_excel", lambda *_a, **_k: ExportReport(alt, ExportKind.XLSX_ALTERNATIVE)
    )
    result = _lote(_write_batch(workspace, [ID_1]))
    assert result.exit_code == 0
    assert "resultado_x_1.xlsx" in result.output


def test_unsafe_path_exits_with_code_2(workspace: Path) -> None:
    result = _lote(_write_batch(workspace, [ID_1]), "--salida", "../fuera")
    assert result.exit_code == 2
    assert "Ruta no permitida" in result.stderr
    assert "../fuera" not in result.stderr
    assert not (workspace / "logs").exists()


def test_salida_subfolder_inside_output_is_used(workspace: Path) -> None:
    result = _lote(_write_batch(workspace, [ID_1]), "--salida", "area")
    assert result.exit_code == 0, result.stderr
    assert list((workspace / "output" / "area").glob("resultado_*.xlsx"))


def test_data_dir_inside_git_tree_exits_2(workspace: Path) -> None:
    (workspace / ".git").mkdir()
    (workspace / ".gitignore").write_text("output/\nlogs/\n", encoding="utf-8")  # sin data/
    result = _lote(_write_batch(workspace, [ID_1]))
    assert result.exit_code == 2
    assert "Ruta no permitida" in result.stderr


def test_log_and_data_dirs_are_checked_with_assert_safe_dir(
    workspace: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    checked: list[Path] = []
    private: list[Path] = []
    real_check, real_private = assert_safe_dir, ensure_private_dir

    def _check(path: Path) -> None:
        checked.append(path)
        real_check(path)

    def _private(path: Path) -> None:
        private.append(path)
        real_private(path)

    monkeypatch.setattr(cli_support, "assert_safe_dir", _check)
    monkeypatch.setattr(cli_support, "ensure_private_dir", _private)
    result = _lote(_write_batch(workspace, [ID_1]))
    assert result.exit_code == 0, result.stderr
    assert Path("data") in checked
    assert Path("logs") in checked
    assert Path("output") in private
    assert Path("logs") in private


def test_input_file_errors_exit_2_with_closed_reason(workspace: Path) -> None:
    bad = workspace / "lote.txt"
    bad.write_text("x", encoding="utf-8")
    result = _lote(bad)
    assert result.exit_code == 2
    assert "formato no soportado" in result.stderr
    missing = _lote(workspace / "no_existe.csv")
    assert missing.exit_code == 2
    assert "no_existe" not in missing.stderr


def test_batch_larger_than_max_exits_2(workspace: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setenv("RUT_MAX_BATCH_SIZE", "1")
    result = _lote(_write_batch(workspace, [ID_1, ID_2]))
    assert result.exit_code == 2
    assert "supera el máximo" in result.stderr


# --- logging, fugas y fallos ---------------------------------------------------------------


def test_build_service_uses_configure_logging(
    workspace: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """L-2: el logger del servicio sale de `configure_logging(run_id, log_dir)`."""
    calls: list[tuple[str, Path]] = []

    def _configure(run_id: str, log_dir: Path) -> logging.Logger:
        calls.append((run_id, log_dir))
        return logging.Logger("sentinel")

    monkeypatch.setattr(cli, "configure_logging", _configure)
    settings = Settings()
    build_service(settings)
    ((run_id, log_dir),) = calls
    assert RUN_ID.fullmatch(run_id)
    assert log_dir == settings.log_dir


def test_run_id_is_shared_by_log_and_excel_and_log_is_closed(workspace: Path) -> None:
    result = _lote(_write_batch(workspace, [ID_1]))
    assert result.exit_code == 0, result.stderr
    (log,) = (workspace / "logs").glob("run_*.jsonl")
    (excel,) = (workspace / "output").glob("resultado_*.xlsx")
    run_id = log.stem.removeprefix("run_")
    assert RUN_ID.fullmatch(run_id)
    assert excel.stem == f"resultado_{run_id}"
    log.unlink()  # en Windows falla si la CLI no cerró el archivo


def test_cli_silences_unconfigured_loggers_at_startup(
    workspace: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    calls: list[str] = []
    monkeypatch.setattr(cli, "silence_unconfigured_loggers", lambda: calls.append("x"))
    assert _lote(_write_batch(workspace, [ID_1])).exit_code == 0
    assert calls == ["x"]


def test_crash_output_has_no_locals_or_full_id(
    workspace: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """Obs-3: un fallo inesperado no vuelca locales ni el texto de la excepción."""
    assert app.pretty_exceptions_show_locals is False

    def _boom(*_a: object, **_k: object) -> list[object]:
        secret_local = ID_1  # noqa: F841  (no debe aparecer en la salida)
        raise RuntimeError(f"fallo con {ID_1}")

    monkeypatch.setattr(cli, "read_batch", _boom)
    result = _lote(_write_batch(workspace, [ID_1]))
    assert result.exit_code == 1
    assert "RuntimeError" in result.stderr
    assert "locals" not in result.stderr.lower()
    assert "Traceback" not in result.stderr
    assert "log" not in result.stderr.lower()  # sin ejecución no hay log al que remitir
    _no_leak(result, workspace, ID_1)


class ChattyPortal(PortalDouble):
    """Portal cuyo código de terceros escribe el id en un logger sin configurar."""

    def query(self, ident: IdentificationInput) -> RutResult:
        logging.getLogger("tercero").warning("consultando %s", ident.identification)
        return super().query(ident)


def test_stderr_never_contains_full_id(workspace: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    portal = ChattyPortal([PortalUnavailable])
    _enable_portal(monkeypatch, portal)
    monkeypatch.setenv("RUT_MAX_RETRIES", "0")
    monkeypatch.setattr(logging.getLogger(), "handlers", [])  # sin pytest: actúa lastResort
    result = _lote(_write_batch(workspace, [ID_1, ID_2, "12A45"]))
    assert result.exit_code == 1
    assert portal.calls == 1  # el tercero sí escribió (fila 1) y no se vio por stderr
    _no_leak(result, workspace, ID_1, ID_2)


def test_output_dir_failure_closes_primary_source_and_leaves_no_log(
    workspace: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    portal = PortalDouble([PORTAL_OK])
    _enable_portal(monkeypatch, portal)
    real_private = ensure_private_dir

    def _private(path: Path) -> None:
        if path == Path("output"):
            raise UnsafePathError(UnsafePathReason.INVALID_PATH)
        real_private(path)

    monkeypatch.setattr(cli_support, "ensure_private_dir", _private)
    result = _lote(_write_batch(workspace, [ID_1]))
    assert result.exit_code == 2
    assert portal.closed == 1
    assert portal.calls == 0
    assert not (workspace / "logs").exists()


def test_unexpected_error_after_run_is_logged_with_class_and_trace_only(
    workspace: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    def _boom(*_a: object, **_k: object) -> ExportReport:
        raise RuntimeError(f"fallo con {ID_1}")

    monkeypatch.setattr(cli, "export_excel", _boom)
    result = _lote(_write_batch(workspace, [ID_1]))
    assert result.exit_code == 1
    assert "RuntimeError" in result.stderr
    assert "log" in result.stderr.lower()  # hay log de la ejecución
    events = _log_events(workspace)
    (event,) = [e for e in events if e["event"] == "exception"]
    assert event["error_code"] == "UNEXPECTED"
    assert event["error_class"] == "RuntimeError"
    trace = event["error_trace"]
    assert isinstance(trace, list)
    assert any(
        str(site).startswith("cli.py:") and str(site).endswith(":_export_and_report")
        for site in trace
    )
    _no_leak(result, workspace, ID_1)


def test_export_error_emits_export_failed_event(
    workspace: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    def _fail(*_a: object, **_k: object) -> ExportReport:
        raise ExportError

    monkeypatch.setattr(cli, "export_excel", _fail)
    assert _lote(_write_batch(workspace, [ID_1])).exit_code == 3
    (event,) = [e for e in _log_events(workspace) if e["event"] == "export_failed"]
    assert event["error_class"] == "ExportError"
    assert event["error_code"] == "EXPORT_FAILED"


def test_json_backup_emits_export_failed_event(
    workspace: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    backup = workspace / "output" / "resultado_x.json"
    monkeypatch.setattr(
        cli, "export_excel", lambda *_a, **_k: ExportReport(backup, ExportKind.JSON_BACKUP)
    )
    assert _lote(_write_batch(workspace, [ID_1])).exit_code == 3
    (event,) = [e for e in _log_events(workspace) if e["event"] == "export_failed"]
    assert event["error_code"] == "JSON_BACKUP"


def test_relative_selectors_path_is_resolved_against_project_root(
    workspace: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    seen: list[Path] = []
    portal = PortalDouble([PORTAL_OK])
    _enable_portal(monkeypatch, portal)

    def _gate(settings: Settings) -> object:
        seen.append(settings.selectors_path)
        return object()

    monkeypatch.setattr(PORTAL_GATE, _gate)
    assert _lote(_write_batch(workspace, [ID_1])).exit_code == 0
    expected = Path(__file__).resolve().parents[2] / "config" / "selectors.yaml"
    assert seen[0].is_absolute()
    assert seen[0].resolve() == expected


def test_output_is_utf8_when_not_a_terminal(tmp_path: Path) -> None:
    (tmp_path / "x.csv").write_text("identificacion\n", encoding="utf-8")
    code = "from rut_checker.cli import app\napp()\n"
    env = {k: v for k, v in os.environ.items() if not k.startswith(("RUT_", "PYTHON"))}
    env["RUT_MAX_RETRIES"] = "zzz"
    completed = subprocess.run(  # noqa: S603  (intérprete propio, código fijo)
        [sys.executable, "-I", "-c", code, "lote", "x.csv"],
        cwd=tmp_path,
        env=env,
        capture_output=True,
        timeout=120,
        check=False,
    )
    assert completed.returncode == 2
    assert "Configuración".encode() in completed.stderr


# --- avisos --------------------------------------------------------------------------------


def test_expired_files_warning(workspace: Path) -> None:
    old = workspace / "output"
    old.mkdir()
    (old / "resultado_20240101-100000-abcd.xlsx").write_text("x", encoding="utf-8")
    result = _lote(_write_batch(workspace, [ID_1]))
    assert result.exit_code == 0, result.stderr
    assert "Hay 1 archivos con más de 365 días: ejecute `rut-checker purgar`" in result.stderr


def test_no_expired_warning_when_nothing_is_old(workspace: Path) -> None:
    result = _lote(_write_batch(workspace, [ID_1]))
    assert "purgar" not in result.stderr


def test_synced_folder_warning(workspace: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(cli_support, "is_synced_folder", lambda _path: True)
    result = _lote(_write_batch(workspace, [ID_1]))
    assert result.exit_code == 0, result.stderr
    assert "sincroniza" in result.stderr.lower()


# --- ayuda ---------------------------------------------------------------------------------


def test_help_lists_commands() -> None:
    result = runner.invoke(app, ["--help"])
    assert result.exit_code == 0
    assert "lote" in result.output
    assert "consultar" in result.output
