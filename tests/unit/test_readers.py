"""Pruebas del lector de lotes `.xlsx` / `.csv` (archivos generados en `tmp_path`)."""

import datetime as dt
import zipfile
import zlib
from pathlib import Path
from typing import NoReturn
from xml.etree.ElementTree import ParseError

import openpyxl
import pytest
from pydantic import ValidationError

from rut_checker.domain.errors import InvalidFileReason, InvalidIdentification, InvalidInputFile
from rut_checker.domain.validators import validate_identification
from rut_checker.io import readers
from rut_checker.io.readers import read_batch


def _write_xlsx(path: Path, rows: list[list[object]]) -> Path:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    assert sheet is not None
    for row in rows:
        sheet.append(row)
    workbook.save(path)
    workbook.close()
    return path


def _write_csv(path: Path, text: str) -> Path:
    path.write_text(text, encoding="utf-8-sig", newline="")
    return path


def _reason(path: Path, max_rows: int = 10) -> InvalidFileReason:
    with pytest.raises(InvalidInputFile) as info:
        read_batch(path, max_rows=max_rows)
    return info.value.reason


def _patch_sheet_xml(path: Path, old: str, new: str) -> None:
    """Edita el XML de la hoja para forzar un valor `<v>` exacto (p. ej. `1000000001.0`)."""
    target = "xl/worksheets/sheet1.xml"
    with zipfile.ZipFile(path) as source:
        members = {name: source.read(name) for name in source.namelist()}
    assert old.encode() in members[target]
    members[target] = members[target].replace(old.encode(), new.encode())
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as out:
        for name, data in members.items():
            out.writestr(name, data)


def test_openpyxl_uses_defusedxml() -> None:
    assert openpyxl.DEFUSEDXML is True


def test_excel_float_keeps_all_digits(tmp_path: Path) -> None:
    """Un `<v>1000000001.0</v>` real llega como float y el dominio lo acepta sin perder dígitos."""
    path = _write_xlsx(tmp_path / "lote.xlsx", [["identificacion"], [1000000001], [1000000002]])
    _patch_sheet_xml(path, "<v>1000000001</v>", "<v>1000000001.0</v>")
    first, second = read_batch(path, max_rows=10)
    assert isinstance(first.raw_identification, float)
    assert isinstance(second.raw_identification, int)
    assert validate_identification(first.raw_identification, None).identification == "1000000001"
    assert validate_identification(second.raw_identification, None).identification == "1000000002"


def test_excel_decimal_float_is_not_read_as_thousands(tmp_path: Path) -> None:
    path = _write_xlsx(tmp_path / "lote.xlsx", [["identificacion"], [1.001]])
    (row,) = read_batch(path, max_rows=10)
    assert row.raw_identification == 1.001
    assert isinstance(row.raw_identification, float)
    with pytest.raises(InvalidIdentification):
        validate_identification(row.raw_identification, None)


@pytest.mark.parametrize(
    ("value", "expected"),
    [
        (1000000001.0, "1000000001"),
        (7.0, "7"),
        (1.001, "1.001"),
        (12, "12"),
        (True, "True"),
        ("  abc ", "abc"),
        ("   ", None),
        (None, None),
        (dt.datetime(2026, 9, 9, 13, 30), "2026-09-09"),
        (dt.date(2026, 9, 10), "2026-09-10"),
        (dt.time(13, 30), "13:30:00"),
        (dt.timedelta(hours=1), "1:00:00"),
    ],
)
def test_cell_to_text_contract(value: object, expected: str | None) -> None:
    assert readers._cell_to_text(value) == expected


def test_excel_date_cell_becomes_iso_string(tmp_path: Path) -> None:
    path = _write_xlsx(
        tmp_path / "lote.xlsx",
        [
            ["identificacion", "fecha_copia_rut"],
            ["1000000001", dt.datetime(2026, 9, 9, 13, 30)],
            ["1000000002", dt.date(2026, 9, 10)],
        ],
    )
    rows = read_batch(path, max_rows=10)
    assert [r.raw_copy_date for r in rows] == ["2026-09-09", "2026-09-10"]


def test_excel_dv_numeric_cell_is_text_without_decimal_zero(tmp_path: Path) -> None:
    path = _write_xlsx(tmp_path / "lote.xlsx", [["identificacion", "dv"], ["900000001", 7]])
    _patch_sheet_xml(path, "<v>7</v>", "<v>7.0</v>")
    (row,) = read_batch(path, max_rows=10)
    assert row.raw_dv == "7"


def test_date_formatted_id_cell_emits_no_warning(
    tmp_path: Path, recwarn: pytest.WarningsRecorder, capsys: pytest.CaptureFixture[str]
) -> None:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.append(["identificacion"])
    sheet.append([1000000001])
    sheet["A2"].number_format = "yyyy-mm-dd"
    path = tmp_path / "lote.xlsx"
    workbook.save(path)
    workbook.close()
    read_batch(path, max_rows=10)
    captured = capsys.readouterr()
    assert len(recwarn) == 0
    assert captured.err == ""
    assert captured.out == ""


def test_reads_copy_columns(tmp_path: Path) -> None:
    path = _write_xlsx(
        tmp_path / "lote.xlsx",
        [
            ["Identificación", "DV", "Fecha_Copia_RUT", "Responsable_Validación", "Observación"],
            ["900000001", 7, "2026-09-09", "USR-001", "nota"],
            ["1000000001", None, None, "", "  "],
        ],
    )
    first, second = read_batch(path, max_rows=10)
    assert first.row_index == 0
    assert first.raw_identification == "900000001"
    assert first.raw_dv == "7"
    assert first.raw_copy_date == "2026-09-09"
    assert first.raw_responsible == "USR-001"
    assert first.raw_note == "nota"
    assert second.row_index == 1
    assert second.raw_dv is None
    assert second.raw_responsible is None
    assert second.raw_note is None


def test_headers_accept_spaces_and_hyphens(tmp_path: Path) -> None:
    path = _write_csv(
        tmp_path / "a.csv",
        "Identificación,Fecha Copia RUT,Responsable-Validación\n1000000001,x,y\n",
    )
    (row,) = read_batch(path, max_rows=10)
    assert (row.raw_copy_date, row.raw_responsible) == ("x", "y")


def test_unknown_columns_are_ignored(tmp_path: Path) -> None:
    path = _write_csv(tmp_path / "a.csv", "nombre,identificacion,otro\nZ,1000000001,w\n")
    (row,) = read_batch(path, max_rows=10)
    assert row.raw_identification == "1000000001"
    assert row.raw_note is None


def test_blank_rows_are_skipped_but_keep_position(tmp_path: Path) -> None:
    path = _write_xlsx(
        tmp_path / "lote.xlsx",
        [["identificacion"], ["1000000001"], [None], ["1000000002"]],
    )
    rows = read_batch(path, max_rows=2)
    assert [(r.row_index, r.raw_identification) for r in rows] == [
        (0, "1000000001"),
        (2, "1000000002"),
    ]


def test_csv_semicolon_and_comma(tmp_path: Path) -> None:
    semicolon = _write_csv(tmp_path / "a.csv", "identificacion;dv\n900000001;7\n1000000001;\n")
    comma = _write_csv(tmp_path / "b.csv", "identificacion,dv\n900000001,7\n1000000001,\n")
    expected = [("900000001", "7"), ("1000000001", None)]
    for path in (semicolon, comma):
        rows = read_batch(path, max_rows=10)
        assert [(r.raw_identification, r.raw_dv) for r in rows] == expected


def test_csv_keeps_unreadable_values_raw(tmp_path: Path) -> None:
    path = _write_csv(tmp_path / "a.csv", "identificacion,fecha_copia_rut\n12A45,ayer\n")
    (row,) = read_batch(path, max_rows=10)
    assert row.raw_identification == "12A45"
    assert row.raw_copy_date == "ayer"


def test_missing_identificacion_column_raises(tmp_path: Path) -> None:
    path = _write_csv(tmp_path / "a.csv", "documento,dv\n1000000001,\n")
    with pytest.raises(InvalidInputFile) as info:
        read_batch(path, max_rows=10)
    assert info.value.reason is InvalidFileReason.MISSING_ID_COLUMN
    assert "1000000001" not in str(info.value)


def test_duplicate_identificacion_column_raises(tmp_path: Path) -> None:
    path = _write_csv(tmp_path / "a.csv", "identificacion,Identificación\n1000000001,1000000002\n")
    assert _reason(path) is InvalidFileReason.DUPLICATE_COLUMN


def test_more_than_max_rows_raises(tmp_path: Path) -> None:
    path = _write_xlsx(
        tmp_path / "lote.xlsx",
        [["identificacion"], ["1000000001"], ["1000000002"], ["1000000003"]],
    )
    assert _reason(path, max_rows=2) is InvalidFileReason.TOO_MANY_ROWS
    assert len(read_batch(path, max_rows=3)) == 3


def test_csv_more_than_max_rows_raises(tmp_path: Path) -> None:
    path = _write_csv(tmp_path / "a.csv", "identificacion\n1000000001\n1000000002\n")
    assert _reason(path, max_rows=1) is InvalidFileReason.TOO_MANY_ROWS


def test_raw_row_cap_stops_reading_blank_padding(tmp_path: Path) -> None:
    padding = "\n" * 30
    path = _write_csv(tmp_path / "a.csv", "identificacion\n1000000001\n" + padding)
    assert _reason(path, max_rows=2) is InvalidFileReason.TOO_MANY_RAW_ROWS


def test_file_larger_than_limit_is_rejected_before_opening(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setattr(readers, "MAX_INPUT_FILE_BYTES", 10)
    path = _write_csv(tmp_path / "a.csv", "identificacion\n1000000001\n")
    assert _reason(path) is InvalidFileReason.TOO_LARGE


@pytest.mark.parametrize("name", ["lote.xls", "lote.xlsm", "lote.txt", "lote"])
def test_unsupported_extension_raises(tmp_path: Path, name: str) -> None:
    path = tmp_path / name
    path.write_bytes(b"x")
    assert _reason(path) is InvalidFileReason.UNSUPPORTED_FORMAT


def test_missing_file_raises(tmp_path: Path) -> None:
    assert _reason(tmp_path / "no_existe.csv") is InvalidFileReason.UNREADABLE


def test_corrupt_xlsx_raises_without_leaking_detail(tmp_path: Path) -> None:
    path = tmp_path / "lote.xlsx"
    path.write_bytes(b"no es un zip")
    with pytest.raises(InvalidInputFile) as info:
        read_batch(path, max_rows=10)
    assert info.value.reason is InvalidFileReason.UNREADABLE
    assert str(path) not in str(info.value)


def test_csv_with_invalid_encoding_raises(tmp_path: Path) -> None:
    path = tmp_path / "a.csv"
    path.write_bytes(b"identificacion\n\xff\xfe\xfa\n")
    assert _reason(path) is InvalidFileReason.NOT_UTF8


def test_empty_csv_raises(tmp_path: Path) -> None:
    assert _reason(_write_csv(tmp_path / "a.csv", "")) is InvalidFileReason.EMPTY


def test_empty_sheet_raises(tmp_path: Path) -> None:
    workbook = openpyxl.Workbook()
    path = tmp_path / "lote.xlsx"
    workbook.save(path)
    workbook.close()
    assert _reason(path) is InvalidFileReason.EMPTY


def test_excel_formula_without_cached_value_is_skipped_keeping_position(tmp_path: Path) -> None:
    # Una fórmula sin valor cacheado (archivo generado por openpyxl) se lee como vacía y la
    # fila se ignora, pero conserva su posición.
    path = _write_xlsx(tmp_path / "lote.xlsx", [["identificacion"], ["=1+1"], ["1000000001"]])
    rows = read_batch(path, max_rows=10)
    assert [(r.row_index, r.raw_identification) for r in rows] == [(1, "1000000001")]


def test_stale_dimension_does_not_drop_rows(tmp_path: Path) -> None:
    path = _write_xlsx(
        tmp_path / "lote.xlsx",
        [["identificacion", "dv"], ["900000001", 7], ["1000000001", None], ["1000000002", None]],
    )
    _patch_sheet_xml(path, '<dimension ref="A1:B4" />', '<dimension ref="A1:A2" />')
    rows = read_batch(path, max_rows=10)
    assert [r.raw_identification for r in rows] == ["900000001", "1000000001", "1000000002"]
    assert rows[0].raw_dv == "7"


def _raise(exc: BaseException) -> object:
    def raiser(*_args: object, **_kwargs: object) -> NoReturn:
        raise exc

    return raiser


@pytest.mark.parametrize(
    "exc",
    [
        NotImplementedError("x"),
        zlib.error("x"),
        ParseError("x"),
        IndexError("x"),
        EOFError("x"),
        zipfile.LargeZipFile("x"),
    ],
)
def test_library_failures_map_to_unreadable(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, exc: BaseException
) -> None:
    path = _write_xlsx(tmp_path / "lote.xlsx", [["identificacion"], ["1000000001"]])
    monkeypatch.setattr(openpyxl, "load_workbook", _raise(exc))
    with pytest.raises(InvalidInputFile) as info:
        read_batch(path, max_rows=10)
    assert info.value.reason is InvalidFileReason.UNREADABLE
    assert info.value.__context__ is None


def test_garbage_sheet_xml_is_unreadable(tmp_path: Path) -> None:
    path = _write_xlsx(tmp_path / "lote.xlsx", [["identificacion"], ["1000000001"]])
    _patch_sheet_xml(path, "<sheetData>", "<sheetData><<<")
    assert _reason(path) is InvalidFileReason.UNREADABLE


def test_truncated_zip_is_unreadable(tmp_path: Path) -> None:
    path = _write_xlsx(tmp_path / "lote.xlsx", [["identificacion"], ["1000000001"]])
    path.write_bytes(path.read_bytes()[:-60])
    assert _reason(path) is InvalidFileReason.UNREADABLE


def test_own_bug_in_row_construction_is_not_reported_as_unreadable(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    path = _write_csv(tmp_path / "a.csv", "identificacion\n1000000001\n")

    def broken(**_kwargs: object) -> NoReturn:
        raise ValidationError.from_exception_data("InputRow", [])

    monkeypatch.setattr(readers, "InputRow", broken)
    with pytest.raises(ValidationError):
        read_batch(path, max_rows=10)


def test_zip_bomb_by_ratio_is_rejected(tmp_path: Path) -> None:
    path = _write_xlsx(tmp_path / "lote.xlsx", [["identificacion"], ["1000000001"]])
    with zipfile.ZipFile(path, "a", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("xl/relleno.bin", b" " * 3_000_000)
    assert _reason(path) is InvalidFileReason.TOO_LARGE


def test_zip_total_uncompressed_size_is_limited(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    path = _write_xlsx(tmp_path / "lote.xlsx", [["identificacion"], ["1000000001"]])
    monkeypatch.setattr(readers, "MAX_UNCOMPRESSED_BYTES", 10)
    assert _reason(path) is InvalidFileReason.TOO_LARGE


def test_very_long_text_is_truncated_so_domain_rejects_it(tmp_path: Path) -> None:
    path = _write_csv(
        tmp_path / "a.csv", "identificacion,observacion\n" + "1" * 5000 + "," + "n" * 5000 + "\n"
    )
    (row,) = read_batch(path, max_rows=10)
    assert isinstance(row.raw_identification, str)
    assert len(row.raw_identification) == readers.MAX_CELL_CHARS + 1
    assert row.raw_note is not None
    assert len(row.raw_note) == readers.MAX_CELL_CHARS + 1


@pytest.mark.parametrize(
    ("value", "expected"),
    [
        ("100\x00000001", "100\ufffd000001"),
        ("a\x7fb", "a\ufffdb"),
        ("a\x1bb", "a\ufffdb"),
        ("a\tb", "a\tb"),
        ("  x  ", "x"),
    ],
)
def test_control_characters_are_replaced_not_removed(value: str, expected: str) -> None:
    assert readers._cell_to_text(value) == expected


def test_nul_in_identification_is_rejected_by_domain_not_merged(tmp_path: Path) -> None:
    path = _write_csv(tmp_path / "a.csv", "identificacion,observacion\n100\x00000001,a\x00b\n")
    (row,) = read_batch(path, max_rows=10)
    assert row.raw_note == "a\ufffdb"
    with pytest.raises(InvalidIdentification):
        validate_identification(row.raw_identification, None)


@pytest.mark.parametrize("column", ["dv", "observacion", "fecha_copia_rut"])
def test_any_duplicated_mapped_column_raises(tmp_path: Path, column: str) -> None:
    path = _write_csv(tmp_path / "a.csv", f"identificacion,{column},{column}\n1000000001,a,b\n")
    assert _reason(path) is InvalidFileReason.DUPLICATE_COLUMN


def test_hidden_first_sheet_is_rejected(tmp_path: Path) -> None:
    workbook = openpyxl.Workbook()
    first = workbook.active
    assert first is not None
    first.append(["identificacion"])
    first.append(["1000000001"])
    workbook.create_sheet("visible")
    first.sheet_state = "hidden"
    workbook.active = 1
    path = tmp_path / "lote.xlsx"
    workbook.save(path)
    workbook.close()
    assert _reason(path) is InvalidFileReason.FIRST_SHEET_HIDDEN


def test_directory_named_like_csv_is_unreadable(tmp_path: Path) -> None:
    folder = tmp_path / "carpeta.csv"
    folder.mkdir()
    assert _reason(folder) is InvalidFileReason.UNREADABLE


@pytest.mark.parametrize("code", [0x80, 0x85, 0x9F, 0xFFFE, 0xFFFF, 0x202A, 0x202E, 0x2066, 0x2069])
def test_c1_noncharacters_and_bidi_controls_are_replaced(code: int) -> None:
    value = "a" + chr(code) + "b"
    assert readers._cell_to_text(value) == "a" + chr(0xFFFD) + "b"


@pytest.mark.parametrize("member", ["xl/sharedStrings.xml", "xl/worksheets/sheet9.xml"])
def test_big_member_with_low_ratio_is_rejected(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, member: str
) -> None:
    path = _write_xlsx(tmp_path / "lote.xlsx", [["identificacion"], ["1000000001"]])
    monkeypatch.setattr(readers, "MAX_CRITICAL_MEMBER_BYTES", 10_000)
    with zipfile.ZipFile(path, "a", zipfile.ZIP_STORED) as archive:
        archive.writestr(member, bytes(range(256)) * 60)
    assert _reason(path) is InvalidFileReason.TOO_LARGE


def test_default_limits_match_policy() -> None:
    assert readers.MAX_UNCOMPRESSED_BYTES == 10 * 1024 * 1024
    assert readers.MAX_CRITICAL_MEMBER_BYTES == 5 * 1024 * 1024
