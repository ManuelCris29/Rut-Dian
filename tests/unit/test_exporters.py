"""Pruebas del exportador Excel/JSON (T-46). Solo identificaciones de la lista blanca."""

import json
from datetime import UTC, datetime, timedelta, timezone
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from rut_checker.domain.errors import ExportError, UnsafePathError, UnsafePathReason
from rut_checker.domain.models import (
    IdentificationInput,
    InputRow,
    QueryOutcome,
    RutResult,
)
from rut_checker.domain.status import QueryStatus, SourceKind
from rut_checker.io import exporters
from rut_checker.io.exporters import (
    COLUMNS,
    ExportKind,
    ExportReport,
    export_excel,
    export_json_backup,
    neutralize_cell,
)

BOGOTA = timezone(timedelta(hours=-5))
WHEN = datetime(2026, 10, 9, 8, 30, tzinfo=BOGOTA)
FORMULA = '=HYPERLINK("http://x","clic")'
ID_ONE = "1000000001"
ID_TWO = "1000000002"
# Valores no incluidos en la lista blanca se arman en tiempo de ejecución.
NOT_WHITELISTED_PHONE = "300" + "1234" + "567"


def _ok(index: int = 0, ident: str = ID_ONE, dv: str | None = None) -> QueryOutcome:
    return QueryOutcome(
        row=InputRow(row_index=index, raw_identification=ident),
        identification=IdentificationInput(identification=ident, dv=dv),
        result=RutResult(
            status=QueryStatus.EXITOSA,
            source=SourceKind.FAKE,
            estado_rut="ACTIVO",
            name="PERSONA FICTICIA UNO",
        ),
        queried_at=WHEN,
        duration_ms=5,
    )


def _invalid(
    raw: str | int | float | None,
    raw_dv: str | None = None,
    raw_note: str | None = None,
    observation: str = "Número inválido: contiene caracteres no numéricos",
) -> QueryOutcome:
    return QueryOutcome(
        row=InputRow(row_index=1, raw_identification=raw, raw_dv=raw_dv, raw_note=raw_note),
        identification=None,
        result=RutResult(
            status=QueryStatus.ENTRADA_INVALIDA,
            source=SourceKind.SIN_CONSULTA,
            observation=observation,
        ),
        queried_at=WHEN,
        duration_ms=0,
    )


def _xlsx(outcomes: list[QueryOutcome], tmp_path: Path, name: str = "r.xlsx") -> Path:
    report = export_excel(outcomes, tmp_path / name, output_dir=tmp_path)
    assert report.kind is ExportKind.XLSX
    return report.path


def _read(path: Path) -> list[list[object]]:
    wb = load_workbook(path)
    assert wb.sheetnames == ["resultado"]
    return [list(r) for r in wb["resultado"].iter_rows(values_only=True)]


def test_one_row_per_input_with_spec_columns(tmp_path: Path) -> None:
    outcomes = [_ok(0), _invalid("12A45"), _ok(2, ID_TWO, dv="3")]
    rows = _read(_xlsx(outcomes, tmp_path, "resultado_x.xlsx"))
    assert tuple(rows[0]) == COLUMNS
    assert list(COLUMNS) == [
        "identificacion",
        "dv",
        "nombre_o_razon_social",
        "estado_rut",
        "fecha_consulta",
        "fuente",
        "estado_consulta",
        "observacion",
    ]
    assert len(rows) == 4
    assert [r[0] for r in rows[1:]] == [ID_ONE, "12A45", ID_TWO]
    assert rows[1][2:4] == ["PERSONA FICTICIA UNO", "ACTIVO"]
    assert rows[3][1] == "3"
    assert rows[2][6] == "ENTRADA_INVALIDA"


def test_every_row_has_source_and_date(tmp_path: Path) -> None:
    rows = _read(_xlsx([_ok(0), _invalid("12A45")], tmp_path))
    for row in rows[1:]:
        assert row[5]
        assert row[4] == "2026-10-09T08:30:00-05:00"
    assert rows[1][5] == "FAKE"
    assert rows[2][5] == "SIN_CONSULTA"


def test_identificacion_written_as_text(tmp_path: Path) -> None:
    outcomes = [_ok(0), _invalid(int(ID_ONE)), _invalid(float(ID_TWO))]
    ws = load_workbook(_xlsx(outcomes, tmp_path))["resultado"]
    for row in (2, 3, 4):
        cell = ws.cell(row=row, column=1)
        assert cell.data_type == "s"
        assert isinstance(cell.value, str)
    assert ws["A3"].value == ID_ONE
    assert ws["A4"].value == ID_TWO


def _outcome_with_text(value: str) -> QueryOutcome:
    return QueryOutcome(
        row=InputRow(row_index=0, raw_identification=ID_ONE),
        identification=IdentificationInput(identification=ID_ONE),
        result=RutResult(
            status=QueryStatus.EXITOSA,
            source=SourceKind.FAKE,
            estado_rut=value,
            name=value,
            observation=value,
        ),
        queried_at=WHEN,
        duration_ms=1,
    )


@pytest.mark.parametrize("prefix", ["=", "+", "-", "@", "\t", "\r"])
def test_formula_like_values_are_neutralized(prefix: str, tmp_path: Path) -> None:
    value = f"{prefix}1+1"
    assert neutralize_cell(value) == "'" + value
    ws = load_workbook(_xlsx([_outcome_with_text(value)], tmp_path))["resultado"]
    for col in (3, 4, 8):
        cell = ws.cell(row=2, column=col)
        assert cell.data_type == "s"
        # El XML normaliza CR a LF al releer; lo relevante es el apostrofo inicial.
        assert cell.value == ("'" + value).replace(chr(13), chr(10))


@pytest.mark.parametrize("prefix", ["=", "+", "-", "@", "\t", "\r"])
def test_json_backup_neutralizes_the_six_prefixes(prefix: str, tmp_path: Path) -> None:
    value = f"{prefix}1+1"
    path = export_json_backup([_outcome_with_text(value)], tmp_path, "resultado_run")
    record = json.loads(path.read_text(encoding="utf-8"))[0]
    for column in ("nombre_o_razon_social", "estado_rut", "observacion"):
        assert record[column] == "'" + value


def test_neutralize_leaves_safe_text_untouched() -> None:
    assert neutralize_cell("ACTIVO") == "ACTIVO"
    assert neutralize_cell("") == ""


def test_formula_in_raw_identification_of_invalid_row_is_neutralized(tmp_path: Path) -> None:
    ws = load_workbook(_xlsx([_invalid(FORMULA, raw_dv=FORMULA)], tmp_path))["resultado"]
    for ref in ("A2", "B2"):
        assert ws[ref].data_type == "s"
        assert ws[ref].value == "'" + FORMULA


ILLEGAL_CHARS = ["\x00", "\x07", "\x1f", "￾", "￿", "\ud800", "\udfff"]


@pytest.mark.parametrize("char", ILLEGAL_CHARS)
def test_control_chars_do_not_break_export_nor_leak(char: str, tmp_path: Path) -> None:
    dirty = ID_ONE[:4] + char + ID_ONE[4:]
    outcome = _invalid(dirty, raw_dv=char, raw_note="a" + char + "b")
    expected = ID_ONE[:4] + "�" + ID_ONE[4:]
    # El .xlsx debe poder releerse (un XML invalido no es un exito).
    ws = load_workbook(_xlsx([outcome], tmp_path))["resultado"]
    assert ws["A2"].value == expected
    assert ws["B2"].value == "�"
    # El JSON debe codificarse y releerse (un surrogate suelto lo romperia).
    path = export_json_backup([outcome], tmp_path, "resultado_run")
    record = json.loads(path.read_text(encoding="utf-8"))[0]
    assert record["identificacion"] == expected
    assert record["dv"] == "�"


def test_illegal_character_error_is_translated_without_text(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setattr(exporters, "_sanitize", lambda value: value)
    dirty = ID_ONE[:4] + "\x00" + ID_ONE[4:]
    report = export_excel([_invalid(dirty)], tmp_path / "r.xlsx", output_dir=tmp_path)
    assert report.kind is ExportKind.JSON_BACKUP
    assert not (tmp_path / "r.xlsx").exists()
    assert not (tmp_path / "r_1.xlsx").exists()


def test_invalid_row_note_is_never_exported(tmp_path: Path) -> None:
    note = f"llamar al {NOT_WHITELISTED_PHONE} hoy"
    outcome = _invalid("12A45", raw_note=note)
    rows = _read(_xlsx([outcome], tmp_path))
    assert rows[1][7] == "Número inválido: contiene caracteres no numéricos"
    path = export_json_backup([outcome], tmp_path, "resultado_run")
    content = path.read_text(encoding="utf-8")
    assert "llamar" not in content
    assert NOT_WHITELISTED_PHONE not in content
    assert "llamar" not in str(rows)


def _lock_workbook_save(monkeypatch: pytest.MonkeyPatch, fail_times: int) -> list[int]:
    original = Workbook.save
    calls: list[int] = []

    def _save(self: Workbook, filename: str | Path) -> None:
        calls.append(1)
        if len(calls) <= fail_times:
            raise PermissionError("bloqueado")
        original(self, filename)

    monkeypatch.setattr(Workbook, "save", _save)
    return calls


def test_locked_file_falls_back_then_json_in_output_dir(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    target = tmp_path / "resultado_run.xlsx"
    calls = _lock_workbook_save(monkeypatch, fail_times=2)
    outcomes = [_ok(0), _invalid("12A45", raw_dv="=1")]
    report = export_excel(outcomes, target, output_dir=tmp_path)
    assert len(calls) == 2
    assert report == ExportReport(path=tmp_path / "resultado_run.json", kind=ExportKind.JSON_BACKUP)
    data = json.loads(report.path.read_text(encoding="utf-8"))
    assert list(data[0]) == list(COLUMNS)
    assert data[0]["identificacion"] == ID_ONE
    assert data[1]["dv"] == "'=1"
    assert sorted(p.name for p in tmp_path.iterdir()) == ["resultado_run.json"]


def test_alternative_name_is_used_when_first_attempt_fails(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    _lock_workbook_save(monkeypatch, fail_times=1)
    report = export_excel([_ok(0)], tmp_path / "resultado_run.xlsx", output_dir=tmp_path)
    assert report == ExportReport(
        path=tmp_path / "resultado_run_1.xlsx", kind=ExportKind.XLSX_ALTERNATIVE
    )
    assert sorted(p.name for p in tmp_path.iterdir()) == ["resultado_run_1.xlsx"]


def test_success_returns_xlsx_kind_and_leaves_no_temp_file(tmp_path: Path) -> None:
    report = export_excel([_ok(0)], tmp_path / "resultado_run.xlsx", output_dir=tmp_path)
    assert report == ExportReport(path=tmp_path / "resultado_run.xlsx", kind=ExportKind.XLSX)
    assert sorted(p.name for p in tmp_path.iterdir()) == ["resultado_run.xlsx"]


def test_all_three_outputs_failing_raises_generic_export_error(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    _lock_workbook_save(monkeypatch, fail_times=2)
    # Un directorio con el nombre del JSON impide el reemplazo atomico.
    (tmp_path / "resultado_run.json").mkdir()
    with pytest.raises(ExportError) as info:
        export_excel([_ok(0)], tmp_path / "resultado_run.xlsx", output_dir=tmp_path)
    assert "bloqueado" not in str(info.value)
    assert ID_ONE not in str(info.value)
    assert not [p for p in tmp_path.iterdir() if ".tmp" in p.name]


def test_target_outside_output_dir_is_rejected(tmp_path: Path) -> None:
    out = tmp_path / "output"
    out.mkdir()
    with pytest.raises(UnsafePathError) as info:
        export_excel([_ok(0)], tmp_path / "otro.xlsx", output_dir=out)
    assert info.value.reason is UnsafePathReason.OUTSIDE_OUTPUT_DIR
    assert not (tmp_path / "otro.xlsx").exists()


def test_wrong_suffix_is_rejected_and_uppercase_is_accepted(tmp_path: Path) -> None:
    with pytest.raises(UnsafePathError) as info:
        export_excel([_ok(0)], tmp_path / "r.txt", output_dir=tmp_path)
    assert info.value.reason is UnsafePathReason.INVALID_PATH
    report = export_excel([_ok(0)], tmp_path / "R.XLSX", output_dir=tmp_path)
    assert report.kind is ExportKind.XLSX


@pytest.mark.parametrize("name", ["", "../x", "a/b", "..", "a\\b"])
def test_json_backup_rejects_unsafe_names(name: str, tmp_path: Path) -> None:
    out = tmp_path / "output"
    out.mkdir()
    with pytest.raises(UnsafePathError):
        export_json_backup([_ok(0)], out, name)
    assert list(out.iterdir()) == []
    assert not (tmp_path / "x.json").exists()


def test_json_backup_columns_only_and_utf8(tmp_path: Path) -> None:
    path = export_json_backup([_ok(0)], tmp_path, "resultado_run")
    assert path == tmp_path / "resultado_run.json"
    data = json.loads(path.read_text(encoding="utf-8"))
    assert data == [
        {
            "identificacion": ID_ONE,
            "dv": "",
            "nombre_o_razon_social": "PERSONA FICTICIA UNO",
            "estado_rut": "ACTIVO",
            "fecha_consulta": "2026-10-09T08:30:00-05:00",
            "fuente": "FAKE",
            "estado_consulta": "EXITOSA",
            "observacion": "",
        }
    ]


def test_json_backup_failure_raises_generic_export_error(tmp_path: Path) -> None:
    missing = tmp_path / "no_existe" / "sub"
    with pytest.raises(ExportError) as info:
        export_json_backup([_ok(0)], missing, "resultado_run")
    assert str(missing) not in str(info.value)


def test_utc_datetime_is_converted_to_bogota(tmp_path: Path) -> None:
    base = _ok(0)
    utc_outcome = QueryOutcome(
        row=base.row,
        identification=base.identification,
        result=base.result,
        queried_at=datetime(2026, 10, 9, 13, 30, tzinfo=UTC),
        duration_ms=1,
    )
    rows = _read(_xlsx([utc_outcome], tmp_path))
    assert rows[1][4] == "2026-10-09T08:30:00-05:00"


def test_existing_unlocked_destination_is_never_overwritten(tmp_path: Path) -> None:
    # Sin bloqueo alguno: un destino existente se trata como fallo y no deja temporales.
    previous = tmp_path / "resultado_run.xlsx"
    previous.write_bytes(b"previo")
    report = export_excel([_ok(0)], previous, output_dir=tmp_path)
    assert previous.read_bytes() == b"previo"
    assert report == ExportReport(
        path=tmp_path / "resultado_run_1.xlsx", kind=ExportKind.XLSX_ALTERNATIVE
    )
    assert not [p for p in tmp_path.iterdir() if ".tmp" in p.name]
